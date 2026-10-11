# -*- coding: utf-8 -*-
"""Pignoraciones de punta a punta: API, bloqueo de ventas, importación del cuadro y patrimonio neto."""
import datetime as dt
import io
import json
import os
import sys
import tempfile
import unittest

_TMP = tempfile.mkdtemp(prefix="rumbo_pg_")
os.environ["PATRIMONIO_DATOS"] = _TMP
os.environ.pop("PATRIMONIO_HOSTS", None)
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook  # noqa: E402

from app import almacen, importar, pignoracion, servidor  # noqa: E402

H = {"Sec-Fetch-Site": "same-origin", "Origin": "http://localhost"}


def cuadro_myinvestor_xlsx(n=120, principal=10000.0, tin=5.0, primero=dt.date(2026, 10, 23)):
    """Réplica de la hoja que descarga MyInvestor (cabecera a partir de la fila 10, columna B)."""
    filas = pignoracion.cuadro({"principal": principal, "interes": tin, "cuotas": n, "fecha": "2026-09-23",
                                "primerPago": primero.isoformat()})
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos MyInvestor"
    ws["D3"], ws["E3"] = "TITULAR:", "NOMBRE APELLIDOS"
    ws["B8"] = "Movimientos"
    for j, t in enumerate(["Fecha", "Cuota", "Intereses", "Amortizado", "Pendiente Amortizar"], start=2):
        ws.cell(row=10, column=j, value=t)
    for i, r in enumerate(filas, start=11):
        f = dt.date.fromisoformat(r["fecha"])
        ws.cell(row=i, column=2, value=f.strftime("%d/%m/%Y"))
        for j, k in enumerate(["cuota", "intereses", "amortizado", "pendiente"], start=3):
            ws.cell(row=i, column=j, value=r[k])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class ConPignoracion(unittest.TestCase):
    def setUp(self):
        for n in os.listdir(_TMP):
            p = os.path.join(_TMP, n)
            if os.path.isfile(p):
                os.remove(p)
        cfg = json.loads(json.dumps(almacen.CARTERA_VACIA))
        cfg["productos"] = [
            {"id": "fondo_a", "nombre": "Fondo A", "corto": "Fondo A", "tipo": "fondo", "fuente": "manual",
             "moneda": "EUR", "slot": 1},
            {"id": "fondo_b", "nombre": "Fondo B", "corto": "Fondo B", "tipo": "fondo", "fuente": "manual",
             "moneda": "EUR", "slot": 2}]
        cfg["valoraciones"] = [
            {"id": "v1", "fecha": "2026-09-01", "producto": "fondo_a", "valor": 3000.0, "aportado": 2500.0},
            {"id": "v2", "fecha": "2026-09-01", "producto": "fondo_b", "valor": 2900.0, "aportado": 2400.0}]
        almacen.guarda(os.path.join(_TMP, "cartera.json"), cfg)
        self.c = servidor.app.test_client()

    def alta(self, **kw):
        d = {"nombre": "Préstamo con mi cartera", "entidad": "MyInvestor", "activos": ["fondo_a", "fondo_b"],
             "principal": "10.000,00", "fecha": "2026-09-23", "interes": "5", "cuotas": 120,
             "primerPago": "2026-10-23"}
        d.update(kw)
        return self.c.post("/api/pignoraciones", json=d, headers=H)

    def test_alta_validacion_y_listado(self):
        r = self.alta()
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        item = r.get_json()["item"]
        self.assertEqual((item["principal"], item["interes"], item["cuotas"], item["estado"]),
                         (10000.0, 5.0, 120, "activa"))
        self.assertEqual(self.alta(principal="", activos=[]).status_code, 400)
        self.assertEqual(self.alta(activos=["no_existe"]).status_code, 400)
        self.assertEqual(self.alta(ltvMax="150").status_code, 400)

    def test_vender_un_activo_pignorado_esta_bloqueado(self):
        self.alta()
        venta = {"producto": "fondo_a", "tipo": "venta", "fecha": "2026-10-01", "importe": "100",
                 "unidades": "0"}
        r = self.c.post("/api/movimientos", json=venta, headers=H)
        self.assertEqual(r.status_code, 400)
        self.assertIn("pignorado", r.get_json()["errores"][0])

    def test_liquidar_libera_y_se_puede_reabrir(self):
        pid = self.alta().get_json()["item"]["id"]
        r = self.c.post(f"/api/pignoraciones/{pid}/liquidar", json={"fecha": "2026-10-05", "importe": "10.050"},
                        headers=H)
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        p = [x for x in r.get_json()["cartera"]["pignoraciones"] if x["id"] == pid][0]
        self.assertEqual((p["estado"], p["fechaLiquidacion"], p["importeLiquidacion"]),
                         ("liquidada", "2026-10-05", 10050.0))
        venta = {"producto": "fondo_a", "tipo": "venta", "fecha": "2026-10-10", "importe": "100", "unidades": "0"}
        self.assertEqual(self.c.post("/api/movimientos", json=venta, headers=H).status_code, 200)
        r = self.c.post(f"/api/pignoraciones/{pid}/reabrir", headers=H)
        self.assertEqual(r.get_json()["item"]["estado"], "activa")

    def test_la_deuda_resta_del_patrimonio(self):
        self.alta()
        datos = servidor.recalcula(descargar=False)
        pg = datos["pignoraciones"][0]
        self.assertGreater(pg["saldo"], 0)
        self.assertEqual(pg["garantia"], 5900.0)
        self.assertEqual([p["tipo"] for p in datos["pasivos"]], ["Deuda / préstamo"])
        self.assertLess(datos["pasivos"][0]["valor"], 0)
        self.assertAlmostEqual(datos["total"]["patrimonio"], 5900.0 + datos["pasivos"][0]["valor"], places=2)
        # los dos fondos salen marcados como pignorados
        self.assertTrue(all(p.get("pignoradoEn") for p in datos["productos"]))

    def test_importar_cuadro_de_myinvestor(self):
        datos, error = importar.leer_cuadro("Amortizaciones_cuenta_1.xlsx", cuadro_myinvestor_xlsx())
        self.assertIsNone(error)
        self.assertEqual((datos["principal"], datos["interes"], datos["cuotas"], datos["cuota"]),
                         (10000.0, 5.0, 120, 106.07))
        self.assertEqual((datos["primerPago"], datos["fechaEstimada"]), ("2026-10-23", "2026-09-23"))
        self.assertAlmostEqual(datos["interesesTotales"], 2727.71, delta=3)

    def test_endpoint_cuadro_y_alta_con_cuadro(self):
        r = self.c.post("/api/pignoraciones/cuadro", headers=H, content_type="multipart/form-data",
                        data={"archivo": (io.BytesIO(cuadro_myinvestor_xlsx()), "cuadro.xlsx")})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        j = r.get_json()
        r = self.alta(cuadro=j["cuadro"], cuotas=None, primerPago=None)
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        self.assertEqual(len(r.get_json()["item"]["cuadro"]), 120)

    def test_archivo_sin_cuadro(self):
        wb = Workbook()
        wb.active["A1"] = "hola"
        b = io.BytesIO()
        wb.save(b)
        datos, error = importar.leer_cuadro("otra.xlsx", b.getvalue())
        self.assertIsNone(datos)
        self.assertIn("cuadro de amortización", error)

    def test_cross_site_no_puede_crear_ni_liquidar(self):
        pid = self.alta().get_json()["item"]["id"]
        mal = {"Sec-Fetch-Site": "cross-site"}
        self.assertEqual(self.c.post("/api/pignoraciones", json={}, headers=mal).status_code, 403)
        self.assertEqual(self.c.post(f"/api/pignoraciones/{pid}/liquidar", json={}, headers=mal).status_code, 403)


if __name__ == "__main__":
    unittest.main()
